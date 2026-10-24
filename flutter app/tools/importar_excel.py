#!/usr/bin/env python3
"""Convierte PLANILLA DE RESERVAS.xlsx en un JSON que la app importa
(Ajustes → Importar planilla).

Fuentes:
  * hoja "diagrama": cada mes es un bloque con una fila por cabaña y una
    columna por día. Una celda pintada es una noche ocupada; el texto va en
    la primera noche ("lorena williman 6p ind.").
  * hojas "desarrollo …": una fila por reserva con montos, contacto, factura
    y canal de publicación.

Las dos se cruzan por cabaña + fecha de llegada: la reserva queda con los
datos de "desarrollo" y, si no está ahí, con lo que dice el diagrama.

Uso:  python3 tools/importar_excel.py "../PLANILLA DE RESERVAS.xlsx" tools/importacion.json
"""
import datetime as dt
import hashlib
import json
import re
import sys
from collections import defaultdict

import openpyxl

CABINS = range(1, 7)
# Antes de julio 2021 el encabezado del diagrama saltea el 22: ahí manda el
# número escrito en la celda y no la posición de la columna.
POSITIONAL_SINCE = dt.date(2021, 7, 1)
WHITE = {'FFFFFFFF', '00000000'}
CHANNELS = [
    ('airbnb', 'airbnb'), ('air bnb', 'airbnb'), ('abnb', 'airbnb'), ('air b&b', 'airbnb'),
    ('booking', 'booking'), ('instagram', 'instagram'), ('facebook', 'facebook'),
    ('google', 'google'), ('web', 'web'), ('recomend', 'recomendado'),
    ('cliente', 'cliente'), ('proveed', 'canje'), ('canje', 'canje'),
]


def fill_key(cell):
    f = cell.fill
    if not f or not f.fill_type:
        return None
    c = f.fgColor
    if c.type == 'rgb':
        return None if c.rgb in WHITE else c.rgb
    if c.type == 'theme':
        if c.theme == 0 and not c.tint:
            return None
        return f't{c.theme}{c.tint:+.1f}' if c.tint else f't{c.theme}'
    return None


def is_green(key):
    # Verdes usados para "ocupado" a lo largo de los años.
    return key is not None and (key.startswith('t9') or key in {'FFA9D18E', 'FFADD58A', 'FFC5E0B4', 'FF92D050', 'FF00B050'})


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def is_block(text):
    t = text.replace(' ', '').lower()
    return len(t) >= 3 and set(t) <= {'x', '.', '-'}


PEOPLE_RE = re.compile(r'(\d+)\s*(?:/\s*(\d+)\s*)?p(?:ers|ax)?\b\.?', re.I)


def parse_label(text):
    """'daiana urdampilleta 8p desa' → nombre, personas, desayuno, mascota, cuna."""
    low = text.lower()
    people = None
    m = PEOPLE_RE.search(text)
    if m:
        people = int(m.group(2) or m.group(1))
    name = PEOPLE_RE.sub(' ', text)
    name = re.sub(r'\b(desay\w*|des\.?|desa|masc\w*|cun\w*|ind\.?|air ?b\w*|abnb|ts|ml)\b', ' ', name, flags=re.I)
    name = re.sub(r'\s+', ' ', name).strip(' .-,')
    return {
        'name': name.title() if name else text.strip(),
        'people': people,
        'breakfast': bool(re.search(r'\bdes(ay\w*|a)?\b', low)),
        'pets': 'masc' in low,
        'crib': 'cun' in low,
    }


def read_diagram(ws):
    """Devuelve las corridas de noches ocupadas por cabaña."""
    months = []  # (fecha_mes, fila_encabezado)
    for row in ws.iter_rows(min_col=1, max_col=1):
        d = as_date(row[0].value)
        if d and d.day == 1:
            months.append((d, row[0].row))

    nights = {}  # (cabaña, fecha) → (texto|None, color)
    for idx, (month, hrow) in enumerate(months):
        end_row = months[idx + 1][1] if idx + 1 < len(months) else ws.max_row
        days_in_month = ((month.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - month).days
        col_day = {}
        for col in range(2, 34):
            if month >= POSITIONAL_SINCE:
                day = col - 1
            else:
                v = ws.cell(hrow, col).value
                day = v if isinstance(v, int) else None
            if day and 1 <= day <= days_in_month:
                col_day[col] = day
        for r in range(hrow + 1, min(end_row, hrow + 9)):
            label = ws.cell(r, 1).value
            m = re.match(r'\s*caba[ñn]a\s*(\d)', str(label or ''), re.I)
            if not m:
                continue
            cabin = int(m.group(1))
            for col, day in col_day.items():
                cell = ws.cell(r, col)
                text = str(cell.value).strip() if cell.value not in (None, '') else None
                key = fill_key(cell)
                if text or key:
                    nights[(cabin, month.replace(day=day))] = (text, key)

    runs, blocks = [], []
    for cabin in CABINS:
        dates = sorted(d for (c, d) in nights if c == cabin)
        current = None
        for d in dates:
            text, key = nights[(cabin, d)]
            if text and is_block(text):
                blocks.append({'cabin': cabin, 'date': d})
                current = None
                continue
            contiguous = current and current['last'] + dt.timedelta(days=1) == d
            if text and not text.isdigit():
                current = {'cabin': cabin, 'checkIn': d, 'last': d, 'label': text, 'key': key}
                runs.append(current)
            elif contiguous and (key == current['key'] or is_green(key)):
                current['last'] = d
            elif is_green(key):
                current = {'cabin': cabin, 'checkIn': d, 'last': d, 'label': None, 'key': key}
                runs.append(current)
            else:
                current = None
    for r in runs:
        r['checkOut'] = r.pop('last') + dt.timedelta(days=1)
        r.pop('key')
    return runs, blocks


def parse_cabins(v):
    if isinstance(v, (int, float)):
        return [int(v)] if 1 <= int(v) <= 6 else []
    return sorted({int(n) for n in re.findall(r'[1-6]', str(v or ''))})


def money(v):
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    return None


def read_sheets(wb):
    rows = []
    for ws in wb.worksheets:
        if not ws.title.lower().startswith(('desarrollo', 'julio', 'mzo')):
            continue
        header = [str(c.value or '').strip().upper() for c in ws[2]]

        def col(*names):
            for n in names:
                for i, h in enumerate(header):
                    if h.startswith(n):
                        return i
            return None

        ix = {
            'name': col('NOMBRE'), 'cabin': col('CABA'), 'in': col('CHECK IN'), 'out': col('CHECK OUT'),
            'total': col('TOTAL'), 'deposit': col('RESERVA'), 'balance': col('SALDO $'),
            'usd': col('SALDO DOLARES', 'DOLARES', 'BOOKING U$'), 'payments': col('PAGOS A CTA'),
            'airbnb': col('AIR BNB'), 'reservedAt': col('FCHA RES'), 'contact': col('CONTACTO'),
            'people': col('PERS', 'CANT. PERS'), 'invoiceDate': col('FACTURA'),
            'invoiceNo': col('FACTURA NRO', 'NUMERO FACT'), 'price': col('VENTA'),
            'nights': col('NOCHES'), 'notes': col('VARIOS'),
        }
        # "FACTURA" (fecha) y "FACTURA NRO." comparten prefijo.
        if ix['invoiceDate'] == ix['invoiceNo']:
            ix['invoiceDate'] = None
        channel_col = len(header) and next((i for i, c in enumerate(ws[1]) if str(c.value or '').upper().startswith('PUBLICACION')), None)

        for row in ws.iter_rows(min_row=3, values_only=True):
            g = lambda k: row[ix[k]] if ix[k] is not None and ix[k] < len(row) else None
            name, cin, cout = g('name'), as_date(g('in')), as_date(g('out'))
            cabins = parse_cabins(g('cabin'))
            if not name or not cin or not cout or not cabins:
                continue
            warnings = []
            nights = g('nights') if isinstance(g('nights'), (int, float)) else None
            if cout <= cin:
                fixed = None
                if nights and len(cabins) == 1:
                    fixed = cout - dt.timedelta(days=int(nights))
                elif cin.replace(year=cin.year - 1) < cout:
                    fixed = cin.replace(year=cin.year - 1)
                if not fixed or fixed >= cout:
                    continue
                warnings.append(f'Fecha de llegada corregida ({cin.isoformat()} → {fixed.isoformat()})')
                cin = fixed
            if (cout - cin).days > 45:
                if not nights or nights > 45:
                    continue
                fixed = cin + dt.timedelta(days=int(nights))
                warnings.append(f'Fecha de salida corregida ({cout.isoformat()} → {fixed.isoformat()})')
                cout = fixed
            channel_raw = ' '.join(str(x) for x in (
                row[channel_col] if channel_col is not None and channel_col < len(row) else None,
                'airbnb' if g('airbnb') not in (None, 0, '') else None, name) if x)
            people = g('people')
            rows.append({
                'sheet': ws.title,
                'name': re.sub(r'\s+', ' ', str(name)).strip().title(),
                'cabins': cabins, 'checkIn': cin, 'checkOut': cout,
                'total': money(g('total')), 'deposit': money(g('deposit')), 'balance': money(g('balance')),
                'balanceUsd': money(g('usd')),
                'reservedAt': as_date(g('reservedAt')),
                'contact': str(g('contact')).strip() if g('contact') not in (None, '') else None,
                'people': int(people) if isinstance(people, (int, float)) else None,
                'invoiceDate': as_date(g('invoiceDate')),
                'invoiceNo': str(g('invoiceNo')).strip() if g('invoiceNo') and not is_block(str(g('invoiceNo'))) else None,
                'nightlyPrice': money(g('price')),
                'priceNote': str(g('price')) if isinstance(g('price'), str) else None,
                'notes': ' · '.join(str(x) for x in (g('notes'), g('payments'), g('usd') if isinstance(g('usd'), str) else None) if x),
                'channel': channel_of(channel_raw),
                'warnings': warnings,
            })
    # Hay hojas que se pisan (mzo-nov20 y 2020-21): una reserva por llegada+cabañas+nombre.
    unique = {}
    for r in rows:
        unique.setdefault((r['name'].lower(), r['checkIn'], tuple(r['cabins'])), r)
    return list(unique.values())


def channel_of(text):
    low = (text or '').lower()
    for needle, ch in CHANNELS:
        if needle in low:
            return ch
    return None


def stable_id(*parts):
    return 'imp-' + hashlib.sha1('|'.join(map(str, parts)).encode()).hexdigest()[:16]


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    runs, blocks = read_diagram(wb['diagrama'])
    wb_values = openpyxl.load_workbook(src, data_only=True)
    sheet_rows = read_sheets(wb_values)

    by_cabin = defaultdict(list)
    for r in runs:
        by_cabin[r['cabin']].append(r)

    def tokens(s):
        return set(re.findall(r'[a-záéíóúñ]{3,}', (s or '').lower()))

    def same_name(label, name):
        # "jar" ↔ "Jarilla", "lorena williman" ↔ "Lorena Williams"
        return any(a.startswith(b) or b.startswith(a) for a in tokens(label) for b in tokens(name))

    def find_run(cabin, row):
        """La corrida del diagrama que corresponde a la fila. Primero una con el mismo
        nombre que se superpone; si no, la que arranca el mismo día (±1 si el
        diagrama no tiene otro nombre escrito)."""
        day = dt.timedelta(days=1)
        best = None
        for r in by_cabin[cabin]:
            if r.get('used') or not (r['checkIn'] < row['checkOut'] and r['checkOut'] > row['checkIn'] - day):
                continue
            diff = abs((r['checkIn'] - row['checkIn']).days)
            named = same_name(r['label'], row['name'])
            if named or diff == 0 or (diff == 1 and not tokens(r['label'])):
                score = (not named, diff)
                if best is None or score < best[0]:
                    best = (score, r)
        return best[1] if best else None

    reservations = []
    matched = 0
    for row in sheet_rows:
        found = {}
        for c in row['cabins']:
            run = find_run(c, row)
            if run:
                run['used'] = True
                found[c] = run
                # Una anotación pegada a la estadía ("camas individuales") que la
                # planilla cuenta dentro de la misma reserva la estira.
                for nxt in by_cabin[c]:
                    if (not nxt.get('used') and nxt['checkIn'] == run['checkOut'] and nxt['label']
                            and nxt['checkOut'] <= row['checkOut'] + dt.timedelta(days=1)
                            and parse_label(nxt['label'])['people'] is None):
                        nxt['used'] = 'nota'
                        run['checkOut'] = nxt['checkOut']
                        row['notes'] = ' · '.join(x for x in (row.get('notes'), nxt['label']) if x)
        if found:
            matched += 1
        # Las cabañas de un grupo pueden tener fechas distintas en el diagrama:
        # una reserva por cada tramo, con la plata en la primera.
        parts = defaultdict(list)
        for c in row['cabins']:
            run = found.get(c)
            parts[(run['checkIn'], run['checkOut']) if run else (row['checkIn'], row['checkOut'])].append(c)
        group_id = stable_id('grupo', row['name'], row['checkIn']) if len(parts) > 1 else None
        for n, ((cin, cout), cabins) in enumerate(sorted(parts.items())):
            runs_here = [found[c] for c in cabins if c in found]
            flags = [parse_label(r['label']) for r in runs_here if r['label']]
            res = {k: v for k, v in row.items() if k != 'sheet'}
            res['warnings'] = list(row['warnings'])
            if n > 0:
                for k in ('total', 'deposit', 'balance', 'balanceUsd', 'nightlyPrice', 'priceNote', 'invoiceNo', 'invoiceDate'):
                    res.pop(k, None)
                res['notes'] = f"Parte de la reserva de {row['name']}: los montos están en la cabaña {sorted(parts.items())[0][1][0]}."
            if (cin, cout) != (row['checkIn'], row['checkOut']):
                res['warnings'].append(
                    f"Fechas tomadas del diagrama (en la planilla: {row['checkIn'].isoformat()} a {row['checkOut'].isoformat()})")
            res.update({
                'id': stable_id(row['name'], row['checkIn'], row['cabins'], *([n] if n else [])),
                'cabins': cabins, 'checkIn': cin, 'checkOut': cout, 'groupId': group_id,
                'source': f"planilla · {row['sheet']}" + (' + diagrama' if runs_here else ''),
                'breakfast': any(f['breakfast'] for f in flags),
                'pets': any(f['pets'] for f in flags),
                'crib': any(f['crib'] for f in flags),
                'diagramLabel': ' / '.join(r['label'] for r in runs_here if r['label']) or None,
            })
            if not res.get('people') or len(parts) > 1:
                res['people'] = sum(f['people'] or 0 for f in flags) or (res.get('people') if n == 0 else None)
            reservations.append(res)

    # Texto suelto dentro de una estadía ya cargada ("camas individuales") es una nota, no otra reserva.
    for run in runs:
        if run.get('used'):
            continue
        for res in reservations:
            if run['cabin'] in res['cabins'] and res['checkIn'] < run['checkIn'] < res['checkOut'] and run['checkOut'] <= res['checkOut'] + dt.timedelta(days=1):
                res['notes'] = ' · '.join(x for x in (res.get('notes'), run['label']) if x)
                run['used'] = 'nota'
                break

    for run in runs:
        if run.get('used'):
            continue
        info = parse_label(run['label']) if run['label'] else {'name': 'Ocupado (sin nombre)', 'people': None, 'breakfast': False, 'pets': False, 'crib': False}
        contact = None
        if not re.search(r'[a-zñ]{2,}', info['name'], re.I):
            contact, info['name'] = run['label'], 'Sin nombre'
        reservations.append({
            'contact': contact,
            'id': stable_id('diagrama', run['cabin'], run['checkIn']),
            'name': info['name'], 'cabins': [run['cabin']],
            'checkIn': run['checkIn'], 'checkOut': run['checkOut'],
            'people': info['people'], 'breakfast': info['breakfast'], 'pets': info['pets'], 'crib': info['crib'],
            'channel': channel_of(run['label']), 'diagramLabel': run['label'],
            'source': 'planilla · diagrama', 'warnings': [] if run['label'] else ['Noches pintadas sin nombre en el diagrama'],
        })

    def ser(v):
        return v.isoformat() if isinstance(v, dt.date) else v

    reservations.sort(key=lambda r: (r['checkIn'], r['cabins']))

    # Superposiciones: se importan igual, marcadas para revisar en la app.
    overlaps = 0
    by_night = defaultdict(list)
    for res in reservations:
        d = res['checkIn']
        while d < res['checkOut']:
            for c in res['cabins']:
                by_night[(c, d)].append(res)
            d += dt.timedelta(days=1)
    for (c, d), lst in by_night.items():
        for res in lst:
            for other in lst:
                if other is not res:
                    msg = f"Se superpone en cabaña {c} con {other['name']} ({other['checkIn'].isoformat()})"
                    res.setdefault('warnings', [])
                    if msg not in res['warnings']:
                        res['warnings'].append(msg)
                        overlaps += 1
    out = {
        'generatedAt': dt.datetime.now().isoformat(timespec='seconds'),
        'source': src.split('/')[-1],
        'stats': {
            'reservations': len(reservations),
            'fromSheets': len(sheet_rows),
            'matchedWithDiagram': matched,
            'onlyInDiagram': sum(1 for r in runs if not r.get('used')),
            'blockedNightsSkipped': len(blocks),
            'withWarnings': sum(1 for r in reservations if r.get('warnings')),
            'overlapWarnings': overlaps,
        },
        'reservations': [{k: ser(v) for k, v in r.items() if v not in (None, '', [])} for r in reservations],
    }
    with open(dst, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(json.dumps(out['stats'], indent=1, ensure_ascii=False))


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
